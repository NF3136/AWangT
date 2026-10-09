#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Read a selected HDX chapter directly into DOCX. No ICS or intermediate HTML.

python export_word.py --chapter "用户手册 > iFC用户手册" --out "Word导出/iFC用户手册.docx"
python export_word.py --list --list-depth 2
Requires: python-docx and lxml (see requirements-word.txt).
"""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import posixpath
import re
import sys
import tempfile
import time
from urllib.parse import quote, unquote, urlsplit
import zipfile
import xml.etree.ElementTree as ET

try:
    from lxml import etree, html
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.image.image import Image
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
except ImportError as exc:
    raise SystemExit("缺少依赖。请执行：python -m pip install -r hdx_export_tools/requirements-word.txt\n" + str(exc))

from export_chapter import Package, PackageError, find_by_path

VERSION = "1.3"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HTML_EXT = {".html", ".htm"}
ATTACH_EXT = {".zip", ".rar", ".7z", ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".csv", ".txt", ".xml"}
BLOCKS = {"div", "p", "section", "article", "main", "blockquote", "dl", "dt", "dd", "figure", "figcaption", "address", "header", "footer"}
IGNORED = {"script", "style", "head", "noscript"}


def clean_text(s):
    return re.sub(r"\s+", "", s or "")


def xml(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def decode(raw, label="未提供文件名"):
    """Strict decoding; never silently replace source text with mojibake."""
    m = re.search(br'charset\s*=\s*["\x27]?([\w-]+)', raw[:4096], re.I)
    enc = m.group(1).decode("ascii") if m else None
    if raw.startswith(b"\xef\xbb\xbf"):
        candidates = ["utf-8-sig"]
    elif raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        candidates = ["utf-32"]
    elif raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        candidates = ["utf-16"]
    else:
        candidates = ([enc] if enc else []) + ["utf-8", "gb18030"]
    failures = []
    for c in dict.fromkeys(candidates):
        try:
            return raw.decode(c), c
        except UnicodeDecodeError as exc:
            failures.append(f"{c}: 字节位置 {exc.start}，原始字节 {raw[exc.start:exc.end].hex(' ')}")
        except LookupError:
            failures.append(f"{c}: Python 不支持该编码名称")
    raise PackageError(
        f"页面编码无法可靠识别：{label}\n声明编码：{enc or '未声明'}\n"
        + "\n".join(failures)
        + "\n已停止；未删除或替换无法识别的字符。请把这段完整错误信息发回排查。")


def bookmark_key(path, fragment=""):
    return "h" + hashlib.sha256((path + "#" + fragment).encode()).hexdigest()[:32]


def resolve_ref(pkg, source, ref):
    u = urlsplit(ref.replace("\\", "/"))
    if u.scheme or u.netloc:
        return None, unquote(u.fragment), True
    p = unquote(u.path)
    target = posixpath.normpath(posixpath.join(posixpath.dirname(source), p)) if p else source
    return pkg.resolve(target), unquote(u.fragment), False


def local_path(root, rel):
    root = Path(root).resolve()
    p = (root / rel).resolve()
    if not p.is_relative_to(root):
        raise PackageError("附件路径越界：" + rel)
    return p


def discover_hdx():
    roots = [Path.cwd(), Path(__file__).resolve().parent.parent]
    found = set()
    for root in roots:
        found.update(p.resolve() for p in root.glob("*.hdx"))
        found.update(p.resolve() for p in root.glob("*/*.hdx"))
    if len(found) != 1:
        raise PackageError(f"发现 {len(found)} 个 HDX，请用 --hdx 指定一个文档包。")
    return next(iter(found))


def select_nodes(pkg, chapters, navi_path="resources/navi.xml"):
    name = pkg.resolve(navi_path)
    if not name:
        raise PackageError("找不到目录索引：" + navi_path)
    nav = ET.fromstring(pkg.read(name))
    nodes = []
    seen_elements = set()

    def visit(el, depth, number):
        if id(el) in seen_elements:
            raise PackageError("选择范围重叠，请不要同时选择父章节和它的子章节。")
        seen_elements.add(id(el))
        url = el.get("url", "")
        target, _, remote = resolve_ref(pkg, name, url)
        if remote or not target:
            raise PackageError(f"目录正文缺失：{el.get('txt')} -> {url}")
        nodes.append({"title": el.get("txt", ""), "path": target, "depth": depth,
                      "number": number, "source_url": url, "id": el.get("id", "")})
        for i, child in enumerate(el.findall("topic"), 1):
            visit(child, depth + 1, number + "." + str(i))

    for i, chapter in enumerate(chapters, 1):
        matches = find_by_path(nav, [p.strip() for p in chapter.split(">")])
        if len(matches) != 1:
            raise PackageError(f"章节匹配数为 {len(matches)}：{chapter}；可先运行 --list。")
        visit(matches[0], 0, str(i))
    return nav, nodes


def read_page(pkg, node):
    # A directory leaf can be an existing Office/PDF file rather than a web page.
    # Keep it byte-for-byte; do not feed binary Office bytes to a text decoder.
    if Path(node["path"]).suffix.lower() not in HTML_EXT:
        return {**node, "kind": "attachment", "encoding": None, "ids": set(),
                "page_ids": [], "first_h1": None, "removed_navigation": []}
    text, encoding = decode(pkg.read(node["path"]), node["path"])
    tree = html.document_fromstring(text, parser=html.HTMLParser(no_network=True))
    candidates = tree.xpath('//*[contains(concat(" ", normalize-space(@class), " "), " articleBoxWithoutHead ") or contains(concat(" ", normalize-space(@class), " "), " articleBox ")]')
    body = candidates[0] if candidates else tree.find("body")
    if body is None:
        raise PackageError("页面没有正文：" + node["path"])
    # Strip website chrome; retain actual lists, inline references, and copyright.
    removed = []
    for e in list(body.iter()):
        if not isinstance(e.tag, str):
            continue
        classes = set(e.get("class", "").split())
        if e.tag in IGNORED or classes & {"ullinks", "footerNavBar", "bottomNavBtn", "breadcrumb"}:
            if e is not body and e.getparent() is not None:
                removed.append(e.tag + ":" + e.get("class", ""))
                e.drop_tree()
    # ICS also removes the per-page mini-TOC in its combined Word export.
    h1s = body.xpath(".//h1")
    first_h1 = h1s[0] if h1s else None
    if first_h1 is not None:
        sibling = first_h1.getnext()
        if sibling is not None and sibling.tag == "ul":
            links = sibling.xpath(".//a[@href]")
            if links and all(a.get("href", "").startswith("#") for a in links):
                if clean_text("".join(sibling.itertext())) == clean_text("".join("".join(a.itertext()) for a in links)):
                    removed.append("page-mini-toc")
                    sibling.drop_tree()
    ids = set()
    for e in body.iter():
        if isinstance(e.tag, str):
            ids.update(v for v in (e.get("id"), e.get("name") if e.tag == "a" else None) if v)
    # Some page-level anchors are outside articleBox; point those to the chapter title.
    page_ids = [e.get("name") for e in tree.xpath("/html/body/a[@name]")]
    return {**node, "kind": "html", "body": body, "encoding": encoding, "ids": ids | set(page_ids),
            "page_ids": page_ids, "first_h1": first_h1, "removed_navigation": removed}


@dataclass
class Flow:
    container: object
    width: int
    p: object = None
    depth: int = 0
    note: bool = False


class WordBuilder:
    def __init__(self, pkg, nodes, out, chapters):
        self.pkg, self.nodes, self.out, self.chapters = pkg, nodes, Path(out), chapters
        self.pages = [read_page(pkg, n) for n in nodes]
        self.paths = {p["path"] for p in self.pages}
        self.ids = {p["path"]: p["ids"] for p in self.pages}
        self.doc = Document()
        self.bookmarks = set()
        self.bookmark_id = 0
        self.outline = []
        self.external = {}
        self.attachments = {}
        self.warnings = []
        self.errors = []
        self.page_reports = []
        self.stats = Counter()
        self.page = None
        self.configure()

    def configure(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
        sec.left_margin, sec.right_margin = Cm(1.4), Cm(1.4)
        sec.header_distance = sec.footer_distance = Cm(.6)
        self.width = int(sec.page_width - sec.left_margin - sec.right_margin)
        styles = self.doc.styles
        for name in ["Normal", "Title", "Subtitle", "Caption"] + [f"Heading {i}" for i in range(1, 10)]:
            s = styles[name]
            s.font.name = "Microsoft YaHei"
            s._element.get_or_add_rPr().append(xml("w:rFonts", eastAsia="微软雅黑", ascii="Calibri", hAnsi="Calibri"))
            s.font.color.rgb = RGBColor(0, 0, 0)
        styles["Normal"].font.size = Pt(10.5)
        styles["Normal"].paragraph_format.space_after = Pt(5)
        styles["Normal"].paragraph_format.line_spacing = 1.2
        for i in range(1, 10):
            s = styles[f"Heading {i}"]
            s.font.size = Pt([18, 15, 13, 12, 11, 11, 11, 11, 11][i - 1])
            s.font.bold = True
            s.paragraph_format.keep_with_next = True
            s.paragraph_format.space_before = Pt(12)
            s.paragraph_format.space_after = Pt(6)
        styles.add_style("HDXGenerated", WD_STYLE_TYPE.CHARACTER)
        code = styles.add_style("HDX Code", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = styles["Normal"]
        code.font.name, code.font.size = "Consolas", Pt(8)
        code.paragraph_format.line_spacing = 1.05
        code.paragraph_format.space_after = Pt(6)
        # Wrap very long technical commands without inserting characters into the text.
        code.element.get_or_add_pPr().append(xml("w:wordWrap", val="0"))
        self.doc.settings.element.append(xml("w:updateFields", val="true"))
        self.doc.core_properties.title = " / ".join(c.split(">")[-1].strip() for c in self.chapters)
        self.doc.core_properties.subject = "从 HDX 指定范围生成的离线手册"

    def bookmark(self, p, key):
        if key in self.bookmarks:
            return
        self.bookmark_id += 1
        p._p.append(xml("w:bookmarkStart", id=self.bookmark_id, name=key))
        p._p.append(xml("w:bookmarkEnd", id=self.bookmark_id))
        self.bookmarks.add(key)

    def paragraph(self, flow):
        if flow.p is None:
            flow.p = flow.container.add_paragraph()
            if flow.depth:
                flow.p.paragraph_format.left_indent = Cm(.45 * flow.depth)
            if flow.note:
                flow.p._p.get_or_add_pPr().append(xml("w:shd", fill="FFF8D9"))
        return flow.p

    def anchors(self, element, flow):
        for v in (element.get("id"), element.get("name") if element.tag == "a" else None):
            if v:
                self.bookmark(self.paragraph(flow), bookmark_key(self.page["path"], v))

    def link_run(self, run, p, link):
        if not link:
            return
        h = xml("w:hyperlink")
        if link[0] == "anchor":
            h.set(qn("w:anchor"), link[1])
        else:
            h.set(qn("r:id"), p.part.relate_to(link[1], RT.HYPERLINK, is_external=True))
        run.font.color.rgb = RGBColor.from_string("0563C1")
        run.font.underline = True
        h.append(run._r)
        p._p.append(h)

    def text(self, text, flow, fmt):
        if not text:
            return
        if not fmt.get("pre"):
            text = re.sub(r"[\t\r\n ]+", " ", text)
            if flow.p is None and not text.strip():
                return
        p = self.paragraph(flow)
        run = p.add_run(text)
        if fmt.get("bold"): run.bold = True
        if fmt.get("italic"): run.italic = True
        if fmt.get("underline"): run.underline = True
        if fmt.get("code") or fmt.get("pre"):
            run.font.name, run.font.size = "Consolas", Pt(8)
        if fmt.get("sup"): run.font.superscript = True
        if fmt.get("sub"): run.font.subscript = True
        self.link_run(run, p, fmt.get("link"))

    def link(self, href):
        actual, fragment, remote = resolve_ref(self.pkg, self.page["path"], href)
        if remote:
            if urlsplit(href).scheme.lower() in {"http", "https", "mailto"}:
                return ("url", href)
            self.warnings.append(f"保留链接文字，未激活不支持的协议：{href}")
            return None
        if actual in self.paths:
            if fragment and fragment not in self.ids[actual]:
                self.errors.append({"source": self.page["path"], "target": href, "reason": "章节内锚点不存在"})
                return self.outside(href, actual, "章节内锚点不存在")
            self.stats["internal_links"] += 1
            return ("anchor", bookmark_key(actual, fragment))
        if actual and Path(actual).suffix.lower() in ATTACH_EXT:
            self.attachments[actual] = self.pkg.read(actual)
            rel = self.out.stem + "_附件/" + actual
            return ("url", quote(rel, safe="/"))
        return self.outside(href, actual, "范围外" if actual else "源包未找到")

    def outside(self, original, actual, reason):
        key = (actual or original, urlsplit(original).fragment, reason)
        if key not in self.external:
            self.external[key] = {"bookmark": "ext" + str(len(self.external) + 1), "target": actual or original,
                                  "original": original, "reason": reason, "sources": []}
        item = self.external[key]
        if self.page["path"] not in item["sources"]:
            item["sources"].append(self.page["path"])
        return ("anchor", item["bookmark"])

    def picture(self, el, flow, fmt):
        ref = el.get("src", "")
        actual, _, remote = resolve_ref(self.pkg, self.page["path"], ref)
        if remote or not actual:
            self.errors.append({"source": self.page["path"], "target": ref, "reason": "图片未在源包找到"})
            self.text("[图片缺失：" + ref + "]", flow, {})
            return
        try:
            data = self.pkg.read(actual)
            info = Image.from_blob(data)
            requested = el.get("width", "")
            px = float(requested) if re.fullmatch(r"\d+(?:\.\d+)?", requested) else info.px_width
            width = min(int(px * 914400 / 96), int(flow.width))
            height_limit = int(Cm(22))
            width = min(width, int(height_limit * info.px_width / info.px_height))
            p = self.paragraph(flow)
            run = p.add_run()
            shape = run.add_picture(BytesIO(data), width=max(width, int(Cm(.2))))
            if el.get("alt"):
                shape._inline.docPr.set("descr", el.get("alt"))
            self.link_run(run, p, fmt.get("link"))
            self.stats["image_occurrences"] += 1
        except Exception as exc:
            self.errors.append({"source": self.page["path"], "target": ref, "reason": "图片转换失败：" + str(exc)})
            self.text("[图片转换失败：" + ref + "]", flow, {})

    def heading(self, text, level, key, container=None, prefix=""):
        p = (container or self.doc).add_paragraph(style=f"Heading {min(9, max(1, level))}")
        if prefix:
            p.add_run(prefix + " ", style="HDXGenerated")
        p.add_run(text)
        self.bookmark(p, key)
        if container is None or container is self.doc:
            self.outline.append((text, level, key, prefix))
        return p

    def children(self, el, flow, fmt=None):
        fmt = fmt or {}
        self.text(el.text, flow, fmt)
        for child in el:
            if isinstance(child.tag, str):
                self.element(child, flow, fmt)
            self.text(child.tail, flow, fmt)

    def element(self, el, flow, fmt):
        tag = el.tag.lower()
        if tag in IGNORED:
            return
        if el is self.page.get("first_h1"):
            # The TOC title represents this heading, but retain its original fragment IDs.
            self.anchors(el, Flow(flow.container, flow.width, self.page_heading))
            return
        if tag == "table":
            flow.p = None
            self.table(el, flow)
            flow.p = None
            return
        if tag in {"ul", "ol"}:
            flow.p = None
            self.anchors(el, flow)
            flow.p = None
            self.text(el.text, flow, fmt)
            num_id = self.numbering(tag == "ol", el.get("start", "1"), el.get("type", "1"))
            for child in el:
                if not isinstance(child.tag, str): continue
                if child.tag == "li":
                    sub = Flow(flow.container, flow.width, depth=flow.depth + 1, note=flow.note)
                    p = self.paragraph(sub)
                    num = xml("w:numPr")
                    num.append(xml("w:ilvl", val="0")); num.append(xml("w:numId", val=num_id))
                    p._p.get_or_add_pPr().append(num)
                    p.paragraph_format.left_indent = Cm(.45 * sub.depth)
                    p.paragraph_format.first_line_indent = Cm(-.35)
                    self.anchors(child, sub)
                    self.children(child, sub, fmt)
                else:
                    self.element(child, flow, fmt)
                self.text(child.tail, flow, fmt)
            flow.p = None
            return
        if re.fullmatch("h[1-6]", tag) or (tag == "p" and "sectiontitle" in el.get("class", "").lower()):
            flow.p = None
            level = self.page["depth"] + (int(tag[1]) if tag.startswith("h") else 3)
            key = bookmark_key(self.page["path"], "_heading_" + str(self.stats["headings"]))
            p = self.heading("".join(el.itertext()).strip(), level, key, flow.container)
            self.stats["headings"] += 1
            self.anchors(el, Flow(flow.container, flow.width, p))
            for child in el.iterdescendants():
                if isinstance(child.tag, str):
                    self.anchors(child, Flow(flow.container, flow.width, p))
            return
        if tag == "pre":
            flow.p = None
            p = flow.container.add_paragraph(style="HDX Code")
            p._p.get_or_add_pPr().append(xml("w:shd", fill="F3F4F5"))
            sub = Flow(flow.container, flow.width, p, flow.depth, flow.note)
            self.anchors(el, sub)
            self.children(el, sub, {**fmt, "pre": True})
            return
        if tag in BLOCKS:
            flow.p = None
            sub = Flow(flow.container, flow.width, depth=flow.depth,
                       note=flow.note or bool(set(el.get("class", "").split()) & {"note", "warning", "notice", "caution", "danger"}))
            self.anchors(el, sub)
            self.children(el, sub, fmt)
            flow.p = None
            return
        self.anchors(el, flow)
        if tag == "img":
            self.picture(el, flow, fmt)
            return
        if tag == "br":
            self.paragraph(flow).add_run().add_break()
            return
        if tag == "hr":
            flow.p = None
            return
        new = dict(fmt)
        if tag in {"b", "strong", "th"}: new["bold"] = True
        if tag in {"i", "em"}: new["italic"] = True
        if tag == "u": new["underline"] = True
        if tag in {"code", "tt", "kbd", "samp"}: new["code"] = True
        if tag == "sup": new["sup"] = True
        if tag == "sub": new["sub"] = True
        if "font-weight:bold" in el.get("style", "").replace(" ", ""): new["bold"] = True
        if tag == "a" and el.get("href"):
            new["link"] = self.link(el.get("href"))
        if tag in {"iframe", "object", "embed", "video", "audio", "svg", "math"}:
            self.warnings.append(f"{self.page['path']} 含特殊元素 {tag}，需要人工核对其文本/媒体。")
        self.children(el, flow, new)

    def numbering(self, ordered, start, kind):
        root = self.doc.part.numbering_part.element
        aid = max([int(e.get(qn("w:abstractNumId"))) for e in root.findall(qn("w:abstractNum"))] + [-1]) + 1
        nid = max([int(e.get(qn("w:numId"))) for e in root.findall(qn("w:num"))] + [0]) + 1
        abstract = xml("w:abstractNum", abstractNumId=aid)
        lvl = xml("w:lvl", ilvl="0")
        lvl.append(xml("w:start", val=int(start) if str(start).isdigit() else 1))
        form = {"a": "lowerLetter", "A": "upperLetter", "i": "lowerRoman", "I": "upperRoman"}.get(kind, "decimal")
        lvl.append(xml("w:numFmt", val=form if ordered else "bullet"))
        lvl.append(xml("w:lvlText", val="%1." if ordered else "•"))
        lvl.append(xml("w:lvlJc", val="left"))
        abstract.append(lvl); root.append(abstract)
        num = xml("w:num", numId=nid); num.append(xml("w:abstractNumId", val=aid)); root.append(num)
        return nid

    def table(self, el, flow):
        caption = el.find("caption")
        if caption is not None:
            sub = Flow(flow.container, flow.width)
            self.children(caption, sub, {"bold": True})
        rows = el.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")
        if not rows: return
        occupied, placements = set(), []
        columns = 0
        for r, row in enumerate(rows):
            c = 0
            for cell in row:
                if cell.tag not in {"td", "th"}: continue
                while (r, c) in occupied: c += 1
                rs = int(cell.get("rowspan", "1")); cs = int(cell.get("colspan", "1"))
                if rs == 0: rs = len(rows) - r
                rs = min(rs, len(rows) - r)
                if not 1 <= rs <= len(rows) or not 1 <= cs <= 100 or c + cs > 100:
                    raise PackageError("不支持的表格跨度：" + self.page["path"])
                box = {(rr, cc) for rr in range(r, r + rs) for cc in range(c, c + cs)}
                if occupied & box:
                    raise PackageError("表格存在重叠合并单元格：" + self.page["path"])
                occupied.update(box); placements.append((r, c, rs, cs, cell)); c += cs
                columns = max(columns, c)
        if not columns: return
        widths = [None] * columns
        for r, c, rs, cs, cell in placements:
            val = cell.get("width", "")
            if val.endswith("%"):
                try:
                    for j in range(c, c + cs):
                        if widths[j] is None: widths[j] = float(val[:-1]) / cs
                except ValueError: pass
        fallback = max(1., (100 - sum(x or 0 for x in widths)) / max(1, widths.count(None)))
        widths = [v or fallback for v in widths]
        widths = [int(flow.width * v / sum(widths)) for v in widths]
        table = flow.container.add_table(rows=len(rows), cols=columns)
        table.style = "Table Grid"
        table.autofit = False
        for i, col in enumerate(table.columns): col.width = widths[i]
        for row in table.rows:
            for i, cell in enumerate(row.cells): cell.width = widths[i]
        self.stats["tables"] += 1
        for r, c, rs, cs, source in placements:
            cell = table.cell(r, c)
            if rs > 1 or cs > 1:
                cell = cell.merge(table.cell(r + rs - 1, c + cs - 1))
                self.stats["merged_cells"] += 1
            for child in list(cell._tc):
                if child.tag != qn("w:tcPr"): cell._tc.remove(child)
            cell.width = sum(widths[c:c + cs])
            sub = Flow(cell, max(int(Cm(.3)), sum(widths[c:c + cs]) - int(Cm(.3))), depth=flow.depth)
            self.anchors(source, sub)
            self.children(source, sub, {"bold": source.tag == "th"})
            if not cell.paragraphs or cell._tc[-1].tag != qn("w:p"):
                cell.add_paragraph()
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.1
            if source.tag == "th":
                cell._tc.get_or_add_tcPr().append(xml("w:shd", fill="EDEFF2"))
        for r, source in enumerate(rows):
            if source.getparent().tag == "thead" or all(c.tag == "th" for c in source if isinstance(c.tag, str)):
                table.rows[r]._tr.get_or_add_trPr().append(xml("w:tblHeader", val="true"))
        # Table-level anchors point to the first cell without making a blank body paragraph.
        self.anchors(el, Flow(table.cell(0, 0), flow.width, table.cell(0, 0).paragraphs[0]))

    def cover(self):
        profile_name = self.pkg.resolve("profile.xml")
        profile = ET.fromstring(self.pkg.read(profile_name)) if profile_name else ET.Element("profile")
        title = self.doc.core_properties.title
        self.doc.add_paragraph(profile.findtext("productType", "") + " " + profile.findtext("productVersion", ""), "Subtitle")
        self.doc.add_paragraph(title, "Title")
        self.doc.add_paragraph(profile.findtext("libName", "产品文档"))
        for label, value in [("文档版本", profile.findtext("libVersion", "")), ("源文档发布日期", profile.findtext("issueDate", "")),
                             ("导出范围", "；".join(self.chapters)), ("导出日期", datetime.now().strftime("%Y-%m-%d"))]:
            if value: self.doc.add_paragraph(label + "：" + value)
        self.doc.add_page_break()
        # Copy copyright text from the source package, never invent legal statements.
        cp = self.pkg.resolve("resources/hedex-copyright.htm")
        if cp:
            text, _ = decode(self.pkg.read(cp), cp)
            t = html.document_fromstring(text)
            body = t.find("body")
            if body is not None:
                for p in body.xpath(".//p | .//h1 | .//h2"):
                    value = "".join(p.itertext()).strip()
                    if value: self.doc.add_paragraph(value)
        else:
            self.doc.add_paragraph("版权所有 © 华为技术有限公司。原始内容及版权归原权利人所有。")
        self.doc.add_page_break()
        self.doc.add_paragraph("目录", "Title")
        self.toc_marker = self.doc.add_paragraph()
        self.doc.add_page_break()
        sec = self.doc.add_section(WD_SECTION_START.CONTINUOUS)
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        sec.header.paragraphs[0].text = profile.findtext("libName", title)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("第 ")
        field = xml("w:fldSimple", instr=" PAGE ")
        r = xml("w:r"); t = xml("w:t"); t.text = "1"; r.append(t); field.append(r); footer._p.append(field)
        footer.add_run(" 页")
        sec._sectPr.append(xml("w:pgNumType", start="1"))

    def toc(self):
        marker = self.toc_marker._p
        # Populate a usable cached TOC now; page references refresh in Word with F9.
        begin = xml("w:r"); begin.append(xml("w:fldChar", fldCharType="begin", dirty="true")); marker.append(begin)
        run = xml("w:r"); instr = xml("w:instrText"); instr.set(qn("xml:space"), "preserve")
        instr.text = ' TOC \\o "1-9" \\h \\z \\u '; run.append(instr); marker.append(run)
        separate = xml("w:r"); separate.append(xml("w:fldChar", fldCharType="separate")); marker.append(separate)
        prev = marker
        for title, depth, key, prefix in self.outline:
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(.35 * min(depth - 1, 8))
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run((prefix + " " if prefix else "") + title)
            self.link_run(r, p, ("anchor", key))
            p.add_run("  ")
            fld = xml("w:fldSimple", instr=" PAGEREF " + key + " \\h ", dirty="true")
            rr = xml("w:r"); tt = xml("w:t"); tt.text = "—"; rr.append(tt); fld.append(rr); p._p.append(fld)
            prev.addnext(p._p); prev = p._p
        end = self.doc.add_paragraph(); rr = xml("w:r"); rr.append(xml("w:fldChar", fldCharType="end")); end._p.append(rr)
        prev.addnext(end._p)

    def body_text(self, elements):
        texts = []
        for el in elements:
            for t in el.iter(qn("w:t")):
                run = t.getparent()
                s = run.find("w:rPr/w:rStyle", self.doc.element.nsmap)
                if s is None or s.get(qn("w:val")) != "HDXGenerated":
                    texts.append(t.text or "")
        return "".join(texts)

    def build(self, on_page=None, on_attachment=None):
        self.cover()
        if on_page:
            # Empty TOC shell only. Each volume builds its own TOC on finalization.
            self.toc()
            self.stream_prefix_count = len(self.doc.element.body) - 1
        used_paths = set()
        for i, page in enumerate(self.pages, 1):
            self.page = page
            print(f"[{i}/{len(self.pages)}] {page['number']} {page['title']}", flush=True)
            if page["kind"] == "attachment" and on_attachment and page["path"] not in used_paths:
                used_paths.add(page["path"])
                self.page_reports.append(on_attachment(self, page))
                continue
            start = len(self.doc.element.body) - 1
            key = bookmark_key(page["path"])
            self.page_heading = self.heading(page["title"], page["depth"] + 1, key, prefix=page["number"])
            if page["path"] in used_paths:
                p = self.doc.add_paragraph("本目录项与前面的正文相同，请点击：")
                self.link_run(p.add_run(page["title"]), p, ("anchor", key))
                self.page_reports.append({k: page[k] for k in ("title", "path", "number", "encoding")} | {"shared_body": True})
                if on_page: on_page(self)
                continue
            used_paths.add(page["path"])
            if page["kind"] == "attachment":
                data = self.pkg.read(page["path"])
                self.attachments[page["path"]] = data
                self.doc.add_paragraph(
                    "此目录项在源文档包中是一份独立文件。原文件已完整保存在随附的附件文件夹中，"
                    "其正文未合并到本手册。请保留 Word 与附件文件夹的相对位置。")
                p = self.doc.add_paragraph()
                relative = self.out.stem + "_附件/" + page["path"]
                self.link_run(p.add_run("打开原文件：" + Path(page["path"]).name), p,
                              ("url", quote(relative, safe="/")))
                self.stats["attachment_nodes"] += 1
                self.warnings.append(f"原文件作为附件保留，未合并正文：{page['title']} -> {page['path']}")
                self.page_reports.append({k: page[k] for k in ("title", "path", "number", "kind", "encoding")} |
                                         {"body_merged": False, "attachment_bytes": len(data),
                                          "attachment_sha256": hashlib.sha256(data).hexdigest(),
                                          "attachment_relative_path": relative})
                if on_page: on_page(self)
                continue
            self.anchors(page["body"], Flow(self.doc, self.width, self.page_heading))
            for aid in page["page_ids"]:
                self.bookmark(self.page_heading, bookmark_key(page["path"], aid))
            before = self.stats.copy()
            self.children(page["body"], Flow(self.doc, self.width))
            elements = list(self.doc.element.body)[start:-1]
            actual = clean_text(self.body_text(elements))
            h1text = "".join(page["first_h1"].itertext()) if page["first_h1"] is not None else ""
            source = "".join(page["body"].itertext())
            # Replace only the first h1 text with the authoritative directory title.
            if h1text:
                source = source.replace(h1text, "", 1)
            expected = clean_text(page["title"] + source)
            intact = expected == actual
            if not intact:
                self.errors.append({"source": page["path"], "reason": "转换前后正文文本不一致", "expected_chars": len(expected), "actual_chars": len(actual)})
            report = {k: page[k] for k in ("title", "path", "number", "kind", "encoding", "removed_navigation")}
            report.update({"text_intact": intact, "source_chars": len(expected), "output_chars": len(actual),
                           "source_text_sha256": hashlib.sha256(expected.encode()).hexdigest(),
                           "output_text_sha256": hashlib.sha256(actual.encode()).hexdigest(),
                           "tables": self.stats["tables"] - before["tables"], "images": self.stats["image_occurrences"] - before["image_occurrences"]})
            self.page_reports.append(report)
            if on_page: on_page(self)
        if self.external:
            self.doc.add_page_break()
            self.doc.add_paragraph("范围外引用及未解析目标", "Title")
            self.doc.add_paragraph("下列目标不在本次正文中，相关链接跳转到本清单。需要对应内容时，请另行选择导出范围。")
            for i, item in enumerate(self.external.values(), 1):
                p = self.doc.add_paragraph(str(i) + ". " + item["target"])
                self.bookmark(p, item["bookmark"])
                self.doc.add_paragraph("状态：" + item["reason"] + "；原引用：" + item["original"])
        if on_page:
            on_page(self, final=True)
            return self.doc
        self.toc()
        # Check generated targets before saving, independent of the source link table.
        names = {e.get(qn("w:name")) for e in self.doc.element.iter(qn("w:bookmarkStart"))}
        for e in self.doc.element.iter(qn("w:hyperlink")):
            anchor = e.get(qn("w:anchor"))
            if anchor and anchor not in names:
                self.errors.append({"target": anchor, "reason": "生成后的 Word 书签缺失"})
        return self.doc


def run_export(hdx, chapters, out, overwrite=False, allow_large=False, dry_run=False):
    start = time.perf_counter()
    out = Path(out).resolve()
    if out.suffix.lower() != ".docx":
        raise PackageError("--out 必须以 .docx 结尾。")
    report_path = out.with_suffix(".report.json")
    if out.exists() and not overwrite:
        raise PackageError(f"文件已存在：{out}\n请换一个输出文件名，或显式加 --overwrite。")
    if report_path.exists() and not overwrite:
        raise PackageError("报告已存在，请换输出名或加 --overwrite：" + str(report_path))
    if out == Path(hdx).resolve():
        raise PackageError("不能覆盖源文件。")
    pkg = Package(str(hdx))
    try:
        _, nodes = select_nodes(pkg, chapters)
        count = len(nodes)
        source_paths = {n["path"] for n in nodes}
        body_count = sum(Path(p).suffix.lower() in HTML_EXT for p in source_paths)
        attachment_count = len(source_paths) - body_count
        print(f"范围：{'；'.join(chapters)}\n目录节点：{count}；HTML 正文：{body_count}；原始文件附件：{attachment_count}\n输出：{out}", flush=True)
        if attachment_count:
            print("注意：原始文件节点将保留在附件文件夹，正文不并入汇总 Word；请将附件与 Word 一起交付。", flush=True)
        if dry_run:
            return {"node_count": count, "unique_source_files": len(source_paths),
                    "unique_body_files": body_count, "attachment_source_files": attachment_count, "output": str(out)}
        if count > 200 and not allow_large:
            raise PackageError(f"本次包含 {count} 个节点，超过默认的 200 项保护阈值。确认范围后加 --allow-large 再运行。")
        builder = WordBuilder(pkg, nodes, out, chapters)
        doc = builder.build()
        out.parent.mkdir(parents=True, exist_ok=True)
        asset_root = out.with_name(out.stem + "_附件")
        if builder.attachments and asset_root.exists() and not overwrite:
            raise PackageError("附件文件夹已存在，请更换输出名或加 --overwrite：" + str(asset_root))
        for name, data in builder.attachments.items():
            target = local_path(asset_root, name)
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(data)
            if hashlib.sha256(target.read_bytes()).digest() != hashlib.sha256(data).digest():
                raise PackageError("附件写入后校验失败：" + name)
        # Save atomically; a failed conversion never truncates an existing Word file.
        with tempfile.NamedTemporaryFile(dir=out.parent, suffix=".docx", delete=False) as f:
            temporary = Path(f.name)
        try:
            doc.save(temporary)
            temporary.replace(out)
        finally:
            temporary.unlink(missing_ok=True)
        report = {"tool_version": VERSION, "source_hdx": str(Path(hdx).resolve()), "chapter_paths": chapters,
                  "output": str(out), "node_count": count, "unique_source_files": len(source_paths),
                  "unique_body_files": body_count, "attachment_source_files": attachment_count,
                  "all_selected_content_merged_into_docx": attachment_count == 0,
                  "pages": builder.page_reports, "counts": dict(builder.stats),
                  "attachments": list(builder.attachments), "outside_references": list(builder.external.values()),
                  "warnings": builder.warnings, "errors": builder.errors,
                  "automatic_checks_passed": not builder.errors, "visual_review": "未执行；请在 Word 中更新目录并检查表格和分页",
                  "elapsed_seconds": round(time.perf_counter() - start, 3), "output_bytes": out.stat().st_size}
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\n已生成：{out}\n报告：{report_path}\n耗时：{report['elapsed_seconds']} 秒；错误：{len(builder.errors)}；提示：{len(builder.warnings)}")
        print("打开 Word 后：Ctrl+A，再按 F9（部分电脑需 Fn+F9），选择更新整个目录。目录中的 — 表示尚未计算页码。")
        if builder.errors:
            print("存在转换错误，文件仅供检查，请先查看报告，不能视为验收通过。")
        return report
    finally:
        pkg.z.close()


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description="从 HDX 直接生成 Word 手册（无需 ICS，无中间 HTML 文件夹）")
    p.add_argument("--hdx", type=Path, help="源 HDX；省略时查找当前目录/工具所在项目中唯一的 HDX")
    p.add_argument("--chapter", action="append", help='章节路径，如 "用户手册 > iFC用户手册"；可重复指定互不重叠的范围')
    p.add_argument("--out", type=Path, help="输出 .docx 路径")
    p.add_argument("--list", action="store_true", help="只列目录，不生成 Word")
    p.add_argument("--list-depth", type=int, default=2, help="列目录层数，默认 2")
    p.add_argument("--dry-run", action="store_true", help="只确认范围和数量，不转换")
    p.add_argument("--overwrite", action="store_true", help="允许覆盖同名输出文件")
    p.add_argument("--allow-large", action="store_true", help="明确允许超过 200 个目录节点的转换")
    p.add_argument("--max-mb", type=float, help="在生成过程中按实际体积分卷，如 10；不生成整本 Word")
    p.add_argument("--split-dir", type=Path, help="分卷保存到一个新文件夹，默认是输出文件名加 _分卷")
    p.add_argument("--split-prefix", help="分卷文件名前缀，例如 用户手册，生成 用户手册1.docx 等")
    p.add_argument("--allow-large-attachments", action="store_true", help="分卷时明确允许原始附件超过 --max-mb 上限")
    p.add_argument("--merge-word-attachments", action="store_true", help="将目录中的 DOC/DOCX 正文按原顺序编入分卷；DOC 需要本机 Microsoft Word")
    args = p.parse_args(argv)
    try:
        if args.max_mb is not None and (not math.isfinite(args.max_mb) or args.max_mb <= 0):
            raise PackageError("--max-mb 必须为正数。")
        hdx = args.hdx or discover_hdx()
        if args.list:
            pkg = Package(str(hdx))
            try:
                root = ET.fromstring(pkg.read(pkg.resolve("resources/navi.xml")))
                def listing(el, path, depth):
                    for child in el.findall("topic"):
                        current = path + [child.get("txt", "")]
                        print(" > ".join(current), f"（{sum(1 for _ in child.iter('topic'))} 项）")
                        if depth < args.list_depth: listing(child, current, depth + 1)
                listing(root, [], 1)
            finally: pkg.z.close()
            return 0
        chapters = args.chapter or ["用户手册 > iFC用户手册"]
        title = chapters[0].split(">")[-1].strip() if len(chapters) == 1 else "所选章节"
        title = re.sub(r'[<>:"/\\|?*]', "_", title)
        out = args.out or Path("Word导出") / (title + ".docx")
        if args.merge_word_attachments and args.max_mb is None:
            raise PackageError("--merge-word-attachments 请同时指定 --max-mb，例如 10。")
        if args.max_mb is not None:
            from stream_word import run_stream_export
            from split_word import SplitError
            try:
                split_dir = args.split_dir or out.with_name(out.stem + "_分卷")
                report = run_stream_export(hdx, chapters, out, split_dir, args.split_prefix or title,
                                           args.max_mb, args.allow_large, args.dry_run,
                                           args.merge_word_attachments, args.allow_large_attachments)
            except SplitError as exc:
                raise PackageError(str(exc)) from exc
        else:
            report = run_export(hdx, chapters, out, args.overwrite, args.allow_large, args.dry_run)
        return 1 if report.get("errors") else 0
    except (PackageError, OSError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
        print("错误：" + str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
