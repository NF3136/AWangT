"""Small fixtures only: streaming, ordering, native attachment content and failure gates."""
from contextlib import redirect_stdout
from io import BytesIO, StringIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from urllib.parse import unquote
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from lxml import etree as E

from export_word import PackageError, WordBuilder, main
from stream_word import run_stream_export
from split_word import W, R, NS, DOC, RELS, SplitError
from test_split_word import noise_png


def fixture(path, attachment=False, huge=False):
    count = 6; images = [noise_png(i, 256 if huge else 128) for i in range(count)]
    native = BytesIO(); doc = Document()
    doc.add_paragraph('附件原生正文，不能遗漏。', 'Title')
    doc.add_paragraph('原生列表', 'List Number')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = '附件合并表格'
    table.cell(1, 0).text = 'A'; table.cell(1, 1).text = 'B'
    doc.add_picture(BytesIO(images[0]), width=Cm(4))
    doc.sections[0].header.paragraphs[0].text = '附件原生页眉'
    doc.save(native)
    with ZipFile(path, 'w') as z:
        child = ''.join(f'<topic txt="章{i}" url="p{i}.html"/>' for i in range(1, count))
        if attachment: child += '<topic txt="附档" url="native.docx"/><topic txt="附档别名" url="native.docx"/>'
        z.writestr('resources/navi.xml', f'<topics><topic txt="手册" url="p0.html">{child}</topic></topics>'.encode())
        for i in range(count):
            h = '手册' if i == 0 else '章' + str(i)
            extra = '<a href="native.docx">附件入口</a>' if attachment and i == 0 else ''
            html = f'<html><meta charset="utf-8"><body><h1>{h}</h1><p>正文段{i}。</p><img src="i{i}.png"><p><a href="p{(i+1)%count}.html">跳转</a>{extra}</p></body></html>'
            z.writestr(f'resources/p{i}.html', html.encode()); z.writestr(f'resources/i{i}.png', images[i])
        if attachment: z.writestr('resources/native.docx', native.getvalue())
    return images, native.getvalue()


class StreamingTests(unittest.TestCase):
    def run_export(self, p, **kwargs):
        with redirect_stdout(StringIO()):
            return run_stream_export(p/'sample.hdx', ['手册'], p/'全集.docx', p/'volumes', '手册', .14, **kwargs)

    def test_streaming_cross_links_and_images(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d); images, _ = fixture(p/'sample.hdx')
            seen_chunks_while_building = []
            original_children = WordBuilder.children
            def observe(builder, *args, **kwargs):
                if builder.page and builder.page['path'].endswith('p5.html'):
                    seen_chunks_while_building.extend(p.glob('.hdx-stream-*/volume-*.docx'))
                return original_children(builder, *args, **kwargs)
            with patch.object(WordBuilder, 'children', observe): report = self.run_export(p)
            self.assertTrue(seen_chunks_while_building, 'Must stage volumes BEFORE all pages finish')
            self.assertFalse((p/'全集.docx').exists()); self.assertFalse(report['whole_book_created'])
            self.assertGreater(len(report['parts']), 1)
            all_text = ''; media = set(); target_names = {}; links = []
            for part in report['parts']:
                file = p/'volumes'/part['file']; self.assertLess(file.stat().st_size, 140_000)
                with ZipFile(file) as z:
                    root = E.fromstring(z.read(DOC)); all_text += ''.join(root.itertext())
                    target_names[file.name] = {e.get('{'+W+'}name') for e in root.iter('{'+W+'}bookmarkStart')}
                    rels = {r.get('Id'): r for r in E.fromstring(z.read(RELS))}
                    for h in root.iter('{'+W+'}hyperlink'):
                        anchor = h.get('{'+W+'}anchor'); rid = h.get('{'+R+'}id')
                        if not anchor: continue
                        if rid: links.append((unquote(rels[rid].get('Target')), anchor))
                        else: self.assertIn(anchor, target_names[file.name])
                    media.update(z.read(n) for n in z.namelist() if n.startswith('word/media/'))
            self.assertEqual(media, set(images)); self.assertTrue(links)
            for file, anchor in links: self.assertIn(anchor, target_names[file])
            for i in range(6): self.assertEqual(all_text.count(f'正文段{i}。'), 1)
            self.assertFalse(list(p.glob('.hdx-stream-*')))

    def test_native_docx_body_and_parts_are_integrated(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d); _, native = fixture(p/'sample.hdx', attachment=True)
            report = self.run_export(p, merge_word_attachments=True)
            self.assertTrue(report['all_selected_content_merged_into_docx'])
            self.assertEqual(report['attachments'], [])
            body = ''; packages = []
            for part in report['parts']:
                with ZipFile(p/'volumes'/part['file']) as z:
                    root = E.fromstring(z.read(DOC)); text = ''.join(root.itertext()); body += text
                    if '附件原生正文' in text: packages.append({n: z.read(n) for n in z.namelist()})
            self.assertEqual(body.count('附件原生正文，不能遗漏。'), 1)
            self.assertEqual(body.count('附件合并表格'), 1)
            self.assertEqual(len(packages), 1)
            with ZipFile(BytesIO(native)) as z:
                for n in ('word/styles.xml', 'word/numbering.xml', 'word/header1.xml'):
                    self.assertEqual(packages[0][n], z.read(n))
            merged = next(v for v in report['pages'] if v.get('body_merged'))
            self.assertEqual(merged['conversion']['converter'], 'DOCX 原始包')

    def test_oversized_indivisible_image_publishes_nothing(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d); fixture(p/'sample.hdx', huge=True)
            with self.assertRaises(SplitError): self.run_export(p)
            self.assertFalse((p/'volumes').exists()); self.assertFalse((p/'全集.docx').exists())

    def test_native_attachment_spans_volumes_with_original_header_and_bookmarks(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d); fixture(p/'sample.hdx', attachment=True)
            doc = Document(); doc.sections[0].header.paragraphs[0].text = '跨卷原生页眉'
            for i in range(5):
                para = doc.add_paragraph('原生段落' + str(i))
                b = OxmlElement('w:bookmarkStart'); b.set(qn('w:id'), str(i)); b.set(qn('w:name'), 'native' + str(i))
                e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), str(i)); para._p.append(b); para._p.append(e)
                doc.add_picture(BytesIO(noise_png(10 + i)), width=Cm(4))
                para = doc.add_paragraph(); h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), 'native' + str((i+1)%5))
                r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '附件内部跳转'; r.append(t); h.append(r); para._p.append(h)
            buf = BytesIO(); doc.save(buf)
            with ZipFile(p/'sample.hdx') as z: members = {n: z.read(n) for n in z.namelist()}
            members['resources/native.docx'] = buf.getvalue()
            with ZipFile(p/'sample.hdx', 'w') as z:
                for name, data in members.items(): z.writestr(name, data)
            report = self.run_export(p, merge_word_attachments=True)
            native_volumes = 0; all_text = ''; names = {}; links = []
            for part in report['parts']:
                with ZipFile(p/'volumes'/part['file']) as z:
                    root = E.fromstring(z.read(DOC)); text = ''.join(root.itertext()); all_text += text
                    names[part['file']] = {e.get(qn('w:name')) for e in root.iter(qn('w:bookmarkStart'))}
                    rels = {r.get('Id'): r for r in E.fromstring(z.read(RELS))}
                    for h in root.iter(qn('w:hyperlink')):
                        rid, anchor = h.get(qn('r:id')), h.get(qn('w:anchor'))
                        if rid and anchor: links.append((unquote(rels[rid].get('Target')), anchor))
                    if '原生段落' in text:
                        native_volumes += 1
                        self.assertIn('跨卷原生页眉', z.read('word/header1.xml').decode())
            self.assertGreater(native_volumes, 1)
            for i in range(5): self.assertEqual(all_text.count('原生段落' + str(i)), 1)
            for file, anchor in links: self.assertIn(anchor, names[file])

    def test_cli_uses_streaming_without_single_book_export(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d); fixture(p/'sample.hdx')
            with patch('export_word.run_export', side_effect=AssertionError('whole-book route used')):
                with redirect_stdout(StringIO()):
                    code = main(['--hdx', str(p/'sample.hdx'), '--chapter', '手册', '--out', str(p/'全集.docx'), '--max-mb', '.14'])
            self.assertEqual(code, 0); self.assertFalse((p/'全集.docx').exists())

    def test_failed_attachment_conversion_never_reports_success(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d); fixture(p/'sample.hdx', attachment=True)
            with patch('stream_word.convert_attachment', side_effect=PackageError('conversion failed')):
                with self.assertRaises(PackageError): self.run_export(p, merge_word_attachments=True)
            self.assertFalse((p/'volumes').exists())

    def test_existing_destination_protected(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d); fixture(p/'sample.hdx'); (p/'volumes').mkdir()
            sentinel = p/'volumes'/'keep'; sentinel.write_text('original')
            with self.assertRaises(PackageError): self.run_export(p)
            self.assertEqual(sentinel.read_text(), 'original')


if __name__ == '__main__': unittest.main()
