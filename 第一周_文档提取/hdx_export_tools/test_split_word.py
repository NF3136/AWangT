# -*- coding: utf-8 -*-
"""Bounded synthetic integration tests for DOCX size splitting."""
from contextlib import redirect_stdout
from io import BytesIO, StringIO
from pathlib import Path
import random
import struct
import tempfile
import unittest
from urllib.parse import unquote
from zipfile import ZipFile
import zlib

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm
from lxml import etree as E

from split_word import split_docx, SplitError, W, R, NS, DOC, RELS, texts, is_chapter


def noise_png(seed, size=128):
    rnd = random.Random(seed)
    raw = b''.join(b'\0' + bytes(rnd.randrange(256) for _ in range(size * 3)) for _ in range(size))
    def chunk(tag, data):
        return struct.pack('>I', len(data)) + tag + data + struct.pack('>I', zlib.crc32(tag + data))
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', size, size, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')


def fixture(path, chapters=6, large_attachment=False, one_image=False):
    doc = Document()
    doc.styles.add_style('HDXGenerated', WD_STYLE_TYPE.CHARACTER)
    doc.add_paragraph('试验手册', 'Title')
    # Exercise existing TOC stripping rather than retaining stale links.
    doc.add_paragraph('目录', 'Title')
    p = doc.add_paragraph()
    for kind in ['begin', 'separate']:
        r = p.add_run()
        if kind == 'separate':
            t = OxmlElement('w:instrText'); t.text = ' TOC \\o "1-9" '; r._r.append(t)
        f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), kind); r._r.append(f)
    doc.add_paragraph('旧目录缓存')
    p = doc.add_paragraph(); f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), 'end'); p.add_run()._r.append(f)
    for i in range(chapters):
        p = doc.add_paragraph(style='Heading 1')
        p.add_run(str(i+1)+' ', style='HDXGenerated'); p.add_run('章节'+str(i+1))
        begin = OxmlElement('w:bookmarkStart'); begin.set(qn('w:id'), str(i)); begin.set(qn('w:name'), 'chapter'+str(i))
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(i)); p._p.append(begin); p._p.append(end)
        doc.add_paragraph('原始正文'+str(i)+'，不可遗漏。')
        doc.add_picture(BytesIO(noise_png(i, 256 if one_image else 128)), width=Cm(5))
        p = doc.add_paragraph(); h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), 'chapter'+str((i+1)%chapters))
        r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '转到下一章'; r.append(t); h.append(r); p._p.append(h)
        table = doc.add_table(rows=2, cols=2); table.cell(0,0).merge(table.cell(0,1)).text = '合并单元格'
        table.cell(1,0).text='A'; table.cell(1,1).text='B'
    if large_attachment:
        attachment=path.parent/'附件/原文件.doc'; attachment.parent.mkdir(); attachment.write_bytes(b'x'*200_000)
        p=doc.add_paragraph(); h=OxmlElement('w:hyperlink')
        h.set(qn('r:id'),doc.part.relate_to('附件/原文件.doc',RT.HYPERLINK,is_external=True))
        r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text='打开附件'; r.append(t); h.append(r); p._p.append(h)
    doc.save(path)


class SplitWordTests(unittest.TestCase):
    def split(self,*args,**kwargs):
        with redirect_stdout(StringIO()): return split_docx(*args,**kwargs)

    def test_actual_size_coverage_images_and_cross_links(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d); source=p/'全集.docx'; fixture(source)
            report=self.split(source,p/'parts',max_mb=.14)
            self.assertGreater(len(report['parts']),1)
            self.assertTrue(report['coverage_contiguous'])
            self.assertTrue(report['all_docx_within_limit'])
            self.assertTrue(report['all_delivered_files_within_limit'])
            cross=0; bodies=[]; targets={}; all_media={}
            with ZipFile(source) as z:
                original=E.fromstring(z.read(DOC)).find('w:body',NS)
                children=list(original); start=next(i for i,n in enumerate(children) if is_chapter(n))
                expected=texts([n for n in children[start:] if n.tag!='{'+W+'}sectPr'])
                source_media={n:z.read(n) for n in z.namelist() if n.startswith('word/media/')}
            for part in report['parts']:
                f=p/'parts'/part['file']; self.assertLessEqual(f.stat().st_size,140_000)
                with ZipFile(f) as z:
                    self.assertIsNone(z.testzip())
                    root=E.fromstring(z.read(DOC)); body=list(root.find('w:body',NS))
                    start=next(i for i,n in enumerate(body) if is_chapter(n))
                    bodies.extend(n for n in body[start:] if n.tag!='{'+W+'}sectPr')
                    self.assertNotIn('旧目录缓存',texts(body))
                    bms={e.get('{'+W+'}name') for e in root.iter('{'+W+'}bookmarkStart')}
                    targets[f.name]=bms
                    rels={r.get('Id'):r for r in E.fromstring(z.read(RELS))}
                    for link in root.iter('{'+W+'}hyperlink'):
                        rid=link.get('{'+R+'}id'); anchor=link.get('{'+W+'}anchor')
                        if rid:
                            if rid.startswith('splitLink'):
                                cross+=1
                                target=unquote(rels[rid].get('Target'))
                                self.assertTrue((p/'parts'/target).exists())
                        elif anchor:self.assertIn(anchor,bms)
                    for n in z.namelist():
                        if n.startswith('word/media/'):
                            self.assertEqual(z.read(n),source_media[n]); all_media[n]=z.read(n)
                    self.assertTrue(root.findall('.//w:gridSpan',NS))
            self.assertEqual(texts(bodies),expected)
            self.assertEqual(all_media,source_media)
            self.assertGreater(cross,0)
            for part in report['parts']:
                for link in part['cross_volume_links']:
                    self.assertIn(link['bookmark'],targets[link['target_file']])

    def test_large_attachment_requires_explicit_exception(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d); source=p/'全集.docx'; fixture(source,chapters=1,large_attachment=True)
            with self.assertRaisesRegex(SplitError,'原始附件'):
                self.split(source,p/'blocked',max_mb=.14)
            self.assertFalse((p/'blocked').exists())
            report=self.split(source,p/'allowed',max_mb=.14,allow_large_attachments=True)
            self.assertFalse(report['all_delivered_files_within_limit'])
            self.assertTrue(report['all_docx_within_limit'])
            self.assertEqual((p/'allowed/附件/原文件.doc').read_bytes(),b'x'*200_000)

    def test_indivisible_oversized_image_fails_without_outputs(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d); fixture(p/'big.docx',chapters=1,one_image=True)
            with self.assertRaisesRegex(SplitError,'单个段落'):
                self.split(p/'big.docx',p/'parts',max_mb=.14)
            self.assertFalse((p/'parts').exists())

    def test_existing_output_directory_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d); fixture(p/'source.docx',chapters=1)
            out=p/'parts';out.mkdir();(out/'keep.txt').write_text('keep')
            with self.assertRaises(SplitError):self.split(p/'source.docx',out)
            self.assertEqual((out/'keep.txt').read_text(),'keep')


if __name__=='__main__':unittest.main(verbosity=2)
